"""Export a Markdown document to a Google Docs-compatible DOCX file.

The exporter creates native headings, paragraphs, bullet lists, code blocks, and
tables. Upload the resulting DOCX to Google Drive and open it with Google Docs.

Usage:
    python tools/export_markdown_to_docx.py docs/feature_catalog.md
    python tools/export_markdown_to_docx.py docs/feature_catalog.md --output docs/feature_catalog.docx
"""

from __future__ import annotations

import argparse
import logging
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Inches, Pt
from markdown_it import MarkdownIt
from markdown_it.token import Token


LOG = logging.getLogger(__name__)


def _plain_text(markdown: str) -> str:
    """Return readable text while retaining the labels of Markdown links."""
    text = re.sub(r"!\[([^]]*)\]\([^)]*\)", r"\1", markdown)
    text = re.sub(r"\[([^]]+)\]\([^)]*\)", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"(\*\*|__|\*|_)", "", text)
    return text.replace("<br>", "\n").strip()


def _configure_document(document: Document) -> None:
    """Set stable page and text styles suitable for wide technical tables."""
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(4)

    code_style = document.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    code_style.font.name = "Consolas"
    code_style.font.size = Pt(8)
    code_style.paragraph_format.space_after = Pt(4)


def _add_table(document: Document, rows: list[list[str]]) -> None:
    """Add a native Word table from parsed Markdown table rows."""
    if not rows:
        return

    n_columns = max(len(row) for row in rows)
    table = document.add_table(rows=0, cols=n_columns)
    table.style = "Table Grid"

    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for column_index, value in enumerate(row):
            cells[column_index].text = _plain_text(value)
            for paragraph in cells[column_index].paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(8)
                    if row_index == 0:
                        run.bold = True


def _collect_table(tokens: list[Token], start_index: int) -> tuple[list[list[str]], int]:
    """Collect table-cell content and return the index after the closing token."""
    rows: list[list[str]] = []
    current_row: list[str] | None = None
    index = start_index + 1

    while index < len(tokens) and tokens[index].type != "table_close":
        token = tokens[index]
        if token.type == "tr_open":
            current_row = []
            rows.append(current_row)
        elif token.type == "inline" and current_row is not None:
            current_row.append(token.content)
        index += 1

    return rows, index + 1


def export_markdown(input_path: Path, output_path: Path, dry_run: bool) -> None:
    """Parse Markdown at input_path and write a native DOCX document."""
    if not input_path.is_file():
        raise FileNotFoundError(f"Markdown input not found: {input_path}")

    markdown = input_path.read_text(encoding="utf-8")
    tokens = MarkdownIt("commonmark").enable("table").parse(markdown)
    document = Document()
    _configure_document(document)

    index = 0
    list_depth = 0
    while index < len(tokens):
        token = tokens[index]
        if token.type == "heading_open":
            content = tokens[index + 1].content
            document.add_heading(_plain_text(content), level=int(token.tag[1]))
            index += 3
            continue
        if token.type == "paragraph_open":
            content = tokens[index + 1].content
            document.add_paragraph(_plain_text(content))
            index += 3
            continue
        if token.type == "bullet_list_open":
            list_depth += 1
        elif token.type == "bullet_list_close":
            list_depth -= 1
        elif token.type == "list_item_open":
            content = tokens[index + 2].content
            style = "List Bullet" if list_depth == 1 else "List Bullet 2"
            document.add_paragraph(_plain_text(content), style=style)
            index += 5
            continue
        elif token.type == "fence":
            document.add_paragraph(token.content.rstrip(), style="Code Block")
        elif token.type == "table_open":
            rows, index = _collect_table(tokens, index)
            _add_table(document, rows)
            continue
        elif token.type == "hr":
            document.add_paragraph("-" * 80)
        index += 1

    if dry_run:
        LOG.info("Dry run: would write %s", output_path)
        return

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    LOG.info("Wrote %s", output_path)


def parse_args() -> argparse.Namespace:
    """Parse command-line arguments."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path, help="Markdown file to export.")
    parser.add_argument(
        "--output",
        type=Path,
        help="DOCX destination. Defaults to the input path with a .docx suffix.",
    )
    parser.add_argument("--dry-run", action="store_true", help="Validate input without writing a DOCX file.")
    parser.add_argument("--verbose", action="store_true", help="Enable debug logging.")
    return parser.parse_args()


def main() -> None:
    """Run the DOCX export command."""
    args = parse_args()
    logging.basicConfig(level=logging.DEBUG if args.verbose else logging.INFO, format="%(levelname)s %(message)s")
    output_path = args.output or args.input.with_suffix(".docx")
    export_markdown(args.input, output_path, args.dry_run)


if __name__ == "__main__":
    main()