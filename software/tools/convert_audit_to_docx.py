"""Convert a Markdown audit document to an editable DOCX file."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.shared import Inches, Pt


SOURCE = Path("docs/session_timing_audit.md")
OUTPUT = Path("docs/session_timing_audit.docx")


def add_inline_text(paragraph, text: str) -> None:
    """Add plain text and Markdown bold runs to a paragraph."""
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        else:
            paragraph.add_run(part)


def add_table(document: Document, rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, value in enumerate(rows[0]):
        cell = table.rows[0].cells[index]
        cell.text = value.strip()
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row in rows[1:]:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value.strip()
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    document.add_paragraph()


def convert(source: Path, output: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)

    lines = source.read_text(encoding="utf-8").splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        if not line.strip():
            index += 1
            continue
        if line.startswith("|") and index + 1 < len(lines) and lines[index + 1].startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].startswith("|"):
                values = [value.strip() for value in lines[index].strip().strip("|").split("|")]
                if not all(set(value) <= {"-", ":", " "} for value in values):
                    table_lines.append(values)
                index += 1
            if table_lines:
                add_table(document, table_lines)
            continue
        heading = re.match(r"^(#{1,6})\s+(.*)$", line)
        if heading:
            paragraph = document.add_paragraph(style=f"Heading {len(heading.group(1))}")
            add_inline_text(paragraph, heading.group(2))
        elif line.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline_text(paragraph, line[2:])
        elif re.match(r"^\d+\.\s+", line):
            paragraph = document.add_paragraph(style="List Number")
            add_inline_text(paragraph, re.sub(r"^\d+\.\s+", "", line))
        else:
            paragraph = document.add_paragraph()
            add_inline_text(paragraph, line)
        index += 1

    document.save(output)


if __name__ == "__main__":
    convert(SOURCE, OUTPUT)